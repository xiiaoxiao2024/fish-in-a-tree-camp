from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
src = ROOT / '每日群文案-How-to-Eat-Fried-Worms.md'
out = ROOT / '每日群文案-How-to-Eat-Fried-Worms.docx'
lines = src.read_text(encoding='utf-8').splitlines()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)
sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.35)
styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); normal.font.size = Pt(10); normal.font.color.rgb = RGBColor(45,52,64)
normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.08
for name,size,color,before,after in [('Heading 1',16,'2E74B5',12,5),('Heading 2',13,'1F4D78',8,4)]:
    st=styles[name]; st.font.name='Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
def shade(p, fill):
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); pPr.append(shd)
title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_after=Pt(3)
r=title.add_run('How to Eat Fried Worms'); r.font.name='Arial'; r.font.size=Pt(24); r.font.bold=True; r.font.color.rgb=RGBColor(0,0,0)
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(12); r=sub.add_run('7天共读营 · 助教每日群发文案'); r.font.name='Arial'; r.font.size=Pt(11); r.font.color.rgb=RGBColor(85,85,85)
for line in lines:
    if not line.strip() or line.startswith('# How') or line.startswith('> 使用说明'):
        continue
    if line.startswith('## '):
        p=doc.add_paragraph(line[3:], style='Heading 1'); p.paragraph_format.page_break_before=True; continue
    if line.startswith('🔗 今日任务卡：'):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.1); p.paragraph_format.space_after=Pt(7); shade(p,'FFF7E6'); p.add_run('今日任务卡：').bold=True; p.add_run(line.split('：',1)[1].strip('`')); continue
    if line.startswith('【') and line.endswith('】'):
        p=doc.add_paragraph(line); p.style='Heading 2'; continue
    p=doc.add_paragraph(line); p.paragraph_format.space_after=Pt(4)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; fr=footer.add_run('How to Eat Fried Worms · 助教工作文案'); fr.font.name='Arial'; fr.font.size=Pt(8); fr.font.color.rgb=RGBColor(120,120,120)
doc.core_properties.title='How to Eat Fried Worms · 7天共读营助教每日群发文案'; doc.core_properties.subject='7天英文共读营'; doc.save(out); print(out)
