from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
src = ROOT / '每日群文案-Tales-of-a-Fourth-Grade-Nothing.md'
out = ROOT / '每日群文案-Tales-of-a-Fourth-Grade-Nothing.docx'
lines = src.read_text(encoding='utf-8').splitlines()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72); sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.85); sec.right_margin = Inches(0.85)

styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(45, 52, 64)
for name, size, color in [('Title', 24, '182848'), ('Heading 1', 17, '3949AB'), ('Heading 2', 12, '3949AB')]:
    st = styles[name]; st.font.name = 'Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(10 if name != 'Title' else 0); st.paragraph_format.space_after = Pt(5)

def shade(paragraph, fill='F3F5FA'):
    pPr = paragraph._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); pPr.append(shd)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Tales of a Fourth Grade Nothing').bold = True
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after = Pt(14)
run = sub.add_run('10 天共读营 · 助教每日群发话术'); run.font.name='Arial'; run.font.size=Pt(11); run.font.color.rgb=RGBColor(95,105,122)

for line in lines:
    if not line.strip() or line.startswith('# Tales') or line.startswith('> 使用说明'):
        continue
    if line.startswith('## Day '):
        if len(doc.paragraphs) > 3: doc.add_page_break()
        p=doc.add_paragraph(line[3:].strip(), style='Heading 1')
        p.paragraph_format.keep_with_next=True
        continue
    if line.startswith('🔗 今日任务卡：'):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.12); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(8); shade(p, 'FFF7E6')
        p.add_run('今日任务卡：').bold=True
        p.add_run(line.split('：',1)[1].strip('`'))
        continue
    p=doc.add_paragraph(line)
    p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr=footer.add_run('Tales of a Fourth Grade Nothing · 助教工作文案'); fr.font.name='Arial'; fr.font.size=Pt(8); fr.font.color.rgb=RGBColor(130,138,150)
doc.core_properties.title = 'Tales of a Fourth Grade Nothing · 助教每日群发话术'
doc.core_properties.subject = '10 天英文共读营'
doc.save(out)
print(out)
